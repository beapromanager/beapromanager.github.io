# Builds the three Word files from PIC/word/{a,b,c}.json (made by scripts/export-conversations.mts)
import json, sys, os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = sys.argv[1]
NAMES = {'a': 'BE A PRO - שיחות - ליגה ג-ב-א', 'b': 'BE A PRO - שיחות - הליגה הלאומית', 'c': 'BE A PRO - שיחות - ליגת העל'}


def rtl(p):
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    pPr.append(bidi)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in p.runs:
        rPr = r._r.get_or_add_rPr()
        r_rtl = OxmlElement('w:rtl')
        rPr.append(r_rtl)
    return p


def para(doc, text, bold=False, size=11, color=None, style=None, indent=0):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = 'Arial'
    run._r.get_or_add_rPr().rFonts.set(qn('w:cs'), 'Arial')
    if color:
        run.font.color.rgb = RGBColor(*color)
    rtl(p)
    if indent:
        p.paragraph_format.right_indent = Pt(indent)
    p.paragraph_format.space_after = Pt(3)
    return p


def heading(doc, text, level):
    p = doc.add_heading(level=level)
    run = p.add_run(text)
    run.font.name = 'Arial'
    run._r.get_or_add_rPr().rFonts.set(qn('w:cs'), 'Arial')
    rtl(p)
    return p


def where(doc, w):
    if w:
        para(doc, '[' + w + ']', bold=True, size=10, color=(0xC0, 0x39, 0x2B))


def money(n):
    return ('+' if n > 0 else '') + format(n, ',') if n else ''


for key, fname in NAMES.items():
    d = json.load(open(os.path.join(SRC, key + '.json'), encoding='utf-8'))
    doc = Document()
    sec = doc.sections[0]
    sec.right_margin = sec.left_margin = Pt(54)
    heading(doc, 'BE A PRO - כל השיחות: ' + d['title'], 0)
    para(doc, 'מסמך לעבור עליו ולאשר. הסכומים מוצגים לפי ' + d['shownAs'] + ('. בליגה ג׳ הם קצת קטנים יותר ובליגה א׳ גדולים בערך פי 3.3.' if key == 'a' else '.'), size=10)
    para(doc, 'סימון באדום בסוגריים מרובעים = הדבר הזה קורה רק בחלק מהליגות של הקבוצה. בלי סימון = קורה בכולן. לתיקון: לכתוב ליד השורה מה לשנות, או למחוק אותה.', size=10)

    heading(doc, '1. מגירת המנהל: דילמות (%d)' % len(d['dilemmas']), 1)
    for i, x in enumerate(d['dilemmas'], 1):
        heading(doc, '%d. %s' % (i, x['id']), 2)
        where(doc, x.get('where'))
        for r in x['rolls']:
            para(doc, r['speaker'] + ': ' + r['text'], bold=True)
            for o in r['options']:
                bits = [b for b in [('כסף ' + money(o['money'])) if o['money'] else '', ('מורל %+d' % o['morale']) if o['morale'] else '', ('אוהדים %+d' % o['fans']) if o['fans'] else ''] if b]
                para(doc, '• ' + o['label'] + ('   (' + ', '.join(bits) + ')' if bits else ''), indent=14)
                para(doc, '→ ' + o['outcome'], size=10, color=(0x55, 0x55, 0x55), indent=30)

    heading(doc, '2. מסיבת עיתונאים (%d שאלות)' % len(d['press']), 1)
    for i, x in enumerate(d['press'], 1):
        where(doc, x.get('where'))
        para(doc, '%d. %s' % (i, x['text']), bold=True)
        for a in x['answers']:
            para(doc, '• ' + a['label'], indent=14)
            para(doc, '→ ' + a['reply'], size=10, color=(0x55, 0x55, 0x55), indent=30)

    heading(doc, '3. הטלפון: צ׳אטים (%d)' % len(d['chats']), 1)
    for i, x in enumerate(d['chats'], 1):
        heading(doc, '%d. %s (%s)' % (i, x['contact'], x['subtitle']), 2)
        where(doc, x.get('where'))
        for l in x['lines']:
            para(doc, l, indent=14)

    heading(doc, '4. החבר בטלפון (%d)' % len(d['mate']), 1)
    for i, x in enumerate(d['mate'], 1):
        heading(doc, '%d. %s' % (i, x['id']), 2)
        where(doc, x.get('where'))
        for l in x['lines']:
            para(doc, l, indent=14)
        for a in x['answers']:
            para(doc, '• ' + a['label'], indent=14)
            para(doc, '→ ' + a['reply'], size=10, color=(0x55, 0x55, 0x55), indent=30)

    heading(doc, '5. היציע: מה האוהדים אומרים (%d שורות)' % len(d['fans']), 1)
    for i, x in enumerate(d['fans'], 1):
        where(doc, x.get('where'))
        para(doc, '%d. [%s] %s' % (i, x['when'], x['text']))

    out = os.path.join(SRC, fname + '.docx')
    doc.save(out)
    print(out)
